"""Excel template and safe preview import for PORTAL invoices."""
import base64
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from portal_excel_workbook import inspect_xlsx

XLSX_MIME='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
MAX_XLSX_BYTES=5*1024*1024
HEADERS=('ID операции','Операция','Количество','Цена, ₽')


def _client_id(value):
    if type(value) is not int or value < 1:
        raise ValueError('Выберите клиента')
    return value


def _operations(service, client_id):
    service.client(client_id, True)
    result=[]
    for op in service.r.catalog('operations'):
        if not op.get('active') or op.get('client_id') != client_id:
            continue
        try:
            tariff=service.tariff(op['id'])
        except ValueError:
            tariff={}
        result.append(dict(
            id=op['id'],
            name=op['name'],
            client_rate=tariff.get('client_rate')
        ))
    return result


def template(service, raw_client_id):
    service.need('invoices.create')
    client_id=_client_id(raw_client_id)
    client=service.client(client_id, True)
    operations=_operations(service, client_id)
    if not operations:
        raise ValueError('У клиента нет активных операций для счёта')

    book=Workbook()
    sheet=book.active
    sheet.title='Счёт'
    sheet.sheet_view.showGridLines=False
    sheet['A1']='PORTAL · Шаблон счёта на оплату'
    sheet['A1'].font=Font(bold=True,size=16)
    sheet['A2']='client_id'
    sheet['B2']=client_id
    sheet.row_dimensions[2].hidden=True
    sheet['A3']='Клиент'
    sheet['B3']=client.get('name','Клиент')
    sheet['A4']='Заполните количество. Цену можно изменить только для этого счёта.'
    sheet.merge_cells('A4:D4')
    sheet['A4'].alignment=Alignment(wrap_text=True)
    for column, value in enumerate(HEADERS, 1):
        cell=sheet.cell(5,column,value)
        cell.font=Font(bold=True,color='FFFFFF')
        cell.fill=PatternFill('solid',fgColor='0B5ED7')
        cell.alignment=Alignment(horizontal='center')
    for row_index, operation in enumerate(operations, 6):
        sheet.cell(row_index,1,operation['id'])
        sheet.cell(row_index,2,operation['name'])
        if operation.get('client_rate') is not None:
            sheet.cell(row_index,4,operation['client_rate']/100)
            sheet.cell(row_index,4).number_format='#,##0.00'
    sheet.column_dimensions['A'].hidden=True
    for col,width in enumerate((14,46,16,18),1):
        sheet.column_dimensions[get_column_letter(col)].width=width
    sheet.freeze_panes='A6'
    sheet.auto_filter.ref=f'A5:D{max(6,5+len(operations))}'
    sheet.protection.sheet=False

    out=BytesIO()
    book.save(out)
    payload=out.getvalue()
    return dict(
        filename=f'PORTAL_invoice_template_client_{client_id}.xlsx',
        mime_type=XLSX_MIME,
        file_b64=base64.b64encode(payload).decode('ascii'),
        client_id=client_id,
        client_name=client.get('name','Клиент')
    )


def _money_to_cents(value):
    if isinstance(value,bool):
        raise ValueError('Цена должна быть числом')
    try:
        number=Decimal(str(value)).quantize(Decimal('0.01'),rounding=ROUND_HALF_UP)
    except (InvalidOperation,ValueError,TypeError):
        raise ValueError('Цена должна быть числом') from None
    if number < 0 or number > Decimal('1000000000'):
        raise ValueError('Цена вне допустимого диапазона')
    return int(number*100)


def preview(service, body):
    service.need('invoices.create')
    if not isinstance(body,dict) or set(body)-{'client_id','file_b64','request_id','company_id'}:
        raise ValueError('Некорректные параметры импорта счёта')
    client_id=_client_id(body.get('client_id'))
    operations={row['id']:row for row in _operations(service,client_id)}
    raw=body.get('file_b64')
    if not isinstance(raw,str) or not raw:
        raise ValueError('Выберите XLSX-файл')
    if len(raw) > ((MAX_XLSX_BYTES+2)//3)*4+8:
        raise ValueError('Файл больше 5 МБ')
    try:
        payload=base64.b64decode(raw,validate=True)
    except (ValueError,TypeError):
        raise ValueError('Некорректный XLSX-файл') from None
    if not 0 < len(payload) <= MAX_XLSX_BYTES:
        raise ValueError('Файл больше 5 МБ')
    inspect_xlsx(payload)
    try:
        book=load_workbook(BytesIO(payload),read_only=False,data_only=False)
    except Exception:
        raise ValueError('Не удалось открыть XLSX-шаблон счёта') from None
    if book.sheetnames != ['Счёт']:
        raise ValueError('Используйте шаблон счёта PORTAL без переименования листа')
    sheet=book['Счёт']
    if sheet['B2'].value != client_id:
        raise ValueError('Шаблон относится к другому клиенту')
    headers=tuple(sheet.cell(5,col).value for col in range(1,5))
    if headers != HEADERS:
        raise ValueError('Заголовки шаблона счёта изменены')
    if sheet.max_row > 505:
        raise ValueError('В счёте слишком много строк')

    lines=[]
    seen=set()
    for row_index in range(6,sheet.max_row+1):
        cells=[sheet.cell(row_index,col) for col in range(1,5)]
        if any(cell.data_type=='f' for cell in cells):
            raise ValueError(f'Формулы запрещены: строка {row_index}')
        operation_id,_,quantity,price=(cell.value for cell in cells)
        if quantity in (None,'') and price in (None,''):
            continue
        try:
            operation_id=int(operation_id)
        except (TypeError,ValueError):
            raise ValueError(f'Строка {row_index}: некорректная операция') from None
        operation=operations.get(operation_id)
        if not operation:
            raise ValueError(f'Строка {row_index}: операция недоступна клиенту')
        if operation_id in seen:
            raise ValueError(f'Строка {row_index}: операция повторяется')
        seen.add(operation_id)
        if isinstance(quantity,bool):
            raise ValueError(f'Строка {row_index}: количество должно быть целым')
        try:
            quantity_int=int(quantity)
        except (TypeError,ValueError):
            raise ValueError(f'Строка {row_index}: количество должно быть целым') from None
        if quantity_int < 1 or quantity_int != quantity:
            raise ValueError(f'Строка {row_index}: количество должно быть больше нуля')
        if price in (None,''):
            if operation.get('client_rate') is None:
                raise ValueError(f'Строка {row_index}: у операции не задана цена клиенту')
            rate=operation['client_rate']
        else:
            rate=_money_to_cents(price)
        lines.append(dict(
            operation_id=operation_id,
            operation_name=operation['name'],
            quantity=quantity_int,
            client_rate=rate,
            amount=quantity_int*rate
        ))
    if not lines:
        raise ValueError('В шаблоне не заполнено количество ни для одной операции')
    return dict(client_id=client_id,lines=lines,total=sum(line['amount'] for line in lines))
