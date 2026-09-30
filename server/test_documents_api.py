"""Synthetic HTTP document permissions, retries, archive and byte-store tests."""
import base64
import hashlib
import json
import sqlite3
import tempfile
import unittest
import sys
import uuid
from datetime import datetime, timedelta
from types import SimpleNamespace
from pathlib import Path
from unittest.mock import patch
from document_domain import LocalFileStorage, Documents, validate_upload, MAX_DOCUMENT_BYTES
from production_repository import Repository
import test_production as fixtures

portal=fixtures.portal
PDF=b'%PDF-1.4\nsynthetic document\n%%EOF\n'

class DocumentAPITest(unittest.TestCase):
    request=fixtures.ProductionTest.request
    tearDown=fixtures.ProductionTest.tearDown
    setUp=fixtures.ProductionTest.setUp
    role_token=fixtures.ProductionTest.role_token
    post=fixtures.ProductionTest.post
    get=fixtures.ProductionTest.get
    work=fixtures.ProductionTest.work

    def upload(self,token=None,**values):
        body=dict(action='upload',document_type='report_pdf',original_filename='Отчёт.pdf',mime_type='application/pdf',
                  file_b64=base64.b64encode(PDF).decode(),title='Отчёт',category='reports',document_date='2026-09-28')
        body.update(values)
        return self.post('documents',body,token)['data']

    def test_upload_storage_checksum_metadata_and_duplicate_request(self):
        a=self.upload(request_id='doc-1');b=self.upload(request_id='doc-1')
        self.assertEqual(a['id'],b['id']);self.assertNotIn('file_b64',a);self.assertNotIn('storage_key',a)
        duplicate=self.upload(request_id='doc-duplicate')
        self.assertEqual(duplicate['id'],a['id'])
        self.assertEqual(a['checksum_sha256'],hashlib.sha256(PDF).hexdigest())
        self.assertEqual(len(self.get('documents')['data']),1)
        download=self.get('document-file?id='+a['id'])['data']
        self.assertEqual(base64.b64decode(download['file_b64']),PDF)
        with portal.db() as conn:
            raw=conn.execute('SELECT metadata,storage_key FROM portal_documents').fetchone()
            self.assertNotIn('file_b64',raw[0]);self.assertTrue(raw[1].startswith('1/'))
        self.post('documents',dict(action='upload',document_type='report_pdf',original_filename='Other.pdf',mime_type='application/pdf',file_b64=base64.b64encode(PDF).decode(),request_id='doc-1'),status=400)

    def test_document_history_returns_visible_revision_chain_and_denies_foreign_company(self):
        first=self.upload(request_id='document-history-v1')
        second=self.upload(request_id='document-history-v2',previous_id=first['id'],title='Отчёт · версия 2')
        self.get('document-history?id='+first['id'],self.worker,status=403)
        versions=self.get('document-history?id='+first['id'])['data']
        self.assertEqual([item['id'] for item in versions],[first['id'],second['id']])
        self.assertEqual([item['revision'] for item in versions],[1,2])
        self.assertNotIn('storage_key',versions[0]);self.assertNotIn('request_id',versions[0])
        other=self.upload(self.other_admin,request_id='foreign-document-history')
        self.get('document-history?id='+other['id'],status=400)
        self.post('documents',dict(action='archive',id=second['id']))
        self.assertEqual(self.get('document-history?id='+second['id'])['data'][-1]['status'],'archived')

    def test_cross_company_read_download_archive_and_reference(self):
        other=self.upload(self.other_admin)
        self.assertEqual(self.get('documents')['data'],[])
        for action in ('document-file','document-metadata'):
            self.get(action+'?id='+other['id'],status=400)
        self.post('documents',dict(action='archive',id=other['id']),status=400)
        self.post('documents',dict(action='upload',document_type='report_pdf',original_filename='x.pdf',mime_type='application/pdf',file_b64=base64.b64encode(PDF).decode(),client_id=900),status=403)
        self.post('documents',dict(company_id=self.other,action='archive',id=other['id']),status=403)

    def test_role_matrix_manager_scope_payroll_and_owner(self):
        doc=self.upload(client_id=1)
        manager=self.role_token('manager');director=self.role_token('director')
        self.assertEqual(self.get('documents',manager)['data'],[])
        with portal.db() as conn:
            conn.execute('INSERT INTO manager_client_assignments(telegram_id,client_id,active,company_id) VALUES(101,1,1,1)')
        self.assertEqual(self.get('documents',manager)['data'][0]['id'],doc['id'])
        self.upload(client_id=None)
        self.assertEqual(len(self.get('documents',manager)['data']),1)
        self.post('documents',dict(action='archive',id=doc['id']),manager,status=403)
        self.assertEqual(len(self.get('documents',director)['data']),2)
        self.get('documents',self.worker,status=403)
        self.request('/api/v3/documents',self.owner,status=403)
        owner=self.request('/api/v3/documents',self.owner,extra_headers={'X-Portal-Company':'1'})
        self.assertEqual(len(owner['data']),2)

    def test_archive_keeps_metadata_blocks_download_and_is_idempotent(self):
        doc=self.upload()
        self.post('documents',dict(action='archive',id=doc['id']))
        self.post('documents',dict(action='archive',id=doc['id']))
        self.assertEqual(self.get('documents')['data'],[])
        self.assertEqual(self.get('documents?status=archived')['data']['total'],1)
        self.get('document-file?id='+doc['id'],status=400)
        self.assertEqual(self.get('document-metadata?id='+doc['id'])['data']['status'],'archived')
        with portal.db() as conn:
            self.assertEqual(conn.execute('SELECT COUNT(*) FROM portal_documents').fetchone()[0],1)
            self.assertEqual(conn.execute("SELECT COUNT(*) FROM portal_production WHERE kind='audit' AND payload LIKE '%document_archived%'").fetchone()[0],1)
            with self.assertRaises(sqlite3.IntegrityError):conn.execute('DELETE FROM portal_documents')

    def test_missing_file_and_corrupt_blob_keep_metadata(self):
        doc=self.upload()
        with portal.db() as conn:key=conn.execute('SELECT storage_key FROM portal_documents').fetchone()[0]
        root=Path(portal.DB_PATH).parent/'.portal-documents'
        (root/key).write_bytes(b'corrupt')
        self.get('document-file?id='+doc['id'],status=400)
        (root/key).unlink()
        self.get('document-file?id='+doc['id'],status=400)
        self.assertEqual(self.get('document-metadata?id='+doc['id'])['data']['id'],doc['id'])

    def test_pagination_search_filters_and_safe_inputs(self):
        self.upload(title='Первый отчёт',client_id=1)
        self.upload(title='Другой документ',category='other',document_date='2026-09-29')
        result=self.get('documents?page=2&limit=1')['data'];self.assertEqual((result['total'],len(result['items'])),(2,1))
        for query in ('q=%D0%9F%D0%B5%D1%80%D0%B2%D1%8B%D0%B9','category=reports','client_id=1','date_from=2026-09-28&date_to=2026-09-28'):
            self.assertEqual(self.get('documents?'+query)['data']['total'],1)
        self.assertEqual(self.get("documents?q=%27%20OR%201%3D1")["data"]["total"],0)
        self.get('documents?limit=201',status=400)
        self.get('documents?date_from=bad',status=400)

    def test_document_employee_filter_is_scoped_and_uses_stable_employee_id(self):
        legacy_employee_id=self.request('/api/me',self.worker)['user']['employee_id']
        with portal.tenants.company_scope(1),portal.db() as conn:
            employee_id=Repository(conn,1).payroll_employee(legacy_employee_id,legacy=True)['employee_id']
        document=self.upload(employee_id=employee_id)
        self.assertEqual(self.get('documents?employee_id='+str(employee_id))['data']['items'][0]['id'],document['id'])
        self.assertEqual(self.get('documents?employee_id=999999')['data']['total'],0)
        self.get('documents?employee_id='+str(employee_id)+'&company_id=2',status=403)

    def test_invalid_filename_mime_size_type_and_employee_reference(self):
        for values in ({'original_filename':'../x.pdf'},{'original_filename':'C:x.pdf'},{'original_filename':'x.pdf '},{'original_filename':' x.pdf'},{'mime_type':'text/plain'},
                       {'document_type':'invoice_xlsx'},{'file_b64':'bad!'},{'file_b64':''},{'employee_id':999999}):
            body=dict(action='upload',document_type='report_pdf',original_filename='x.pdf',mime_type='application/pdf',file_b64=base64.b64encode(PDF).decode())
            body.update(values);self.post('documents',body,status=400)
        with self.assertRaises(ValueError):validate_upload('x.pdf','application/pdf',b'x'*(MAX_DOCUMENT_BYTES+1),'report_pdf')
        self.assertEqual(self.get('documents')['data'],[])

    def test_invoice_pdf_generation_permissions_tenant_scope_and_financial_immutability(self):
        work=self.work()
        invoice=self.post('invoices',dict(work_ids=[work['id']]))['data']
        before=(self.get('works')['data'],self.get('invoices')['data'])
        renderer=SimpleNamespace(invoice_pdf=lambda *args: PDF)
        import documents_api
        errors=[];original_route=documents_api.route
        def traced(*args):
            try:return original_route(*args)
            except Exception as error:errors.append(error);raise
        with patch.dict(sys.modules,{'pdf_documents':renderer}),patch('documents_api.route',side_effect=traced):
            try:doc=self.post('document-generate',dict(document_type='invoice_pdf',invoice_id=invoice['id']))['data']
            except AssertionError:self.fail('PDF route failed: '+repr(errors))
            self.assertEqual(doc['document_type'],'invoice_pdf')
            self.assertEqual(doc['client_id'],invoice['client_id'])
            self.assertEqual(base64.b64decode(self.get('document-file?id='+doc['id'])['data']['file_b64']),PDF)
            self.post('document-generate',dict(document_type='invoice_pdf',invoice_id=invoice['id']),self.worker,status=403)
            self.post('document-generate',dict(document_type='invoice_pdf',invoice_id=invoice['id']),self.other_admin,status=400)
            self.post('document-generate',dict(document_type='payroll_slip_pdf',payroll_period_id='unknown',employee_id=1),self.worker,status=403)
        after=(self.get('works')['data'],self.get('invoices')['data'])
        self.assertEqual(after,before)
        with portal.db() as conn:
            row=conn.execute('SELECT status,document_type FROM portal_documents WHERE id=?',(doc['id'],)).fetchone()
            self.assertEqual(tuple(row),('ready','invoice_pdf'))

    def test_invoice_editing_revision_permissions_and_payment_lock(self):
        work=self.work();invoice=self.post('invoices',dict(work_ids=[work['id']]))['data']
        self.assertEqual((invoice['state'],invoice['revision']),('finalized',1))
        before_works=self.get('works')['data']
        renderer=SimpleNamespace(invoice_pdf=lambda *args: PDF)
        with patch.dict(sys.modules,{'pdf_documents':renderer}):
            old_document=self.post('document-generate',dict(document_type='invoice_pdf',invoice_id=invoice['id']))['data']
        manager=self.role_token('manager')
        self.post('invoices',dict(workflow='send_to_editing',invoice_id=invoice['id']),manager,status=403)
        self.post('invoices',dict(workflow='send_to_editing',invoice_id=invoice['id']),self.other_admin,status=400)
        self.request('/api/v3/invoices',self.admin,dict(workflow='send_to_editing',invoice_id=invoice['id'],request_id=str(uuid.uuid4())),extra_headers={'X-Portal-Company':str(self.other)},status=403)
        editing=self.post('invoices',dict(workflow='send_to_editing',invoice_id=invoice['id']))['data']
        self.assertEqual(editing['state'],'editing')
        repeated=self.post('invoices',dict(workflow='send_to_editing',invoice_id=invoice['id']))['data']
        self.assertEqual((repeated['state'],repeated['revision']),('editing',editing['revision']))
        with portal.db() as conn:conn.execute('INSERT INTO manager_client_assignments(telegram_id,client_id,active,company_id) VALUES(101,1,1,1)')
        original_line=invoice['lines'][0]
        updated=self.post('invoices',dict(workflow='save_revision',invoice_id=invoice['id'],lines=[dict(work_id=work['id'],quantity=1,client_rate=original_line['client_rate']//2)],due_at='2026-10-01T00:00:00'),manager)['data']
        self.assertEqual((updated['state'],updated['revision']),('finalized',2))
        self.assertEqual(updated['history'][0]['revision'],1)
        self.assertNotEqual(updated['amount'],invoice['amount'])
        self.assertEqual((updated['lines'][0]['quantity'],updated['lines'][0]['client_rate']),(1,original_line['client_rate']//2))
        self.post('invoices',dict(workflow='save_revision',invoice_id=invoice['id'],work_ids=[work['id']]),manager,status=400)
        with patch.dict(sys.modules,{'pdf_documents':renderer}):
            new_document=self.post('document-generate',dict(document_type='invoice_pdf',invoice_id=invoice['id']))['data']
        self.assertNotEqual(old_document['id'],new_document['id'])
        self.assertEqual(self.get('document-metadata?id='+old_document['id'])['data']['status'],'ready')
        self.assertEqual(self.get('works')['data'],before_works)
        work2=self.work();work3=self.work()
        invoice2=self.post('invoices',dict(work_ids=[work2['id'],work3['id']]))['data']
        self.post('payments',dict(invoice_id=invoice2['id'],amount=str(invoice2['amount']/100)))
        self.post('invoices',dict(workflow='send_to_editing',invoice_id=invoice2['id']))
        self.post('invoices',dict(workflow='save_revision',invoice_id=invoice2['id'],work_ids=[work2['id']]),status=400)

    def test_invoice_xlsx_is_valid_printable_and_idempotent(self):
        from io import BytesIO
        from openpyxl import load_workbook
        with portal.db() as conn:
            requisites=('ООО Портал','7700000000','770001001','ОГРН 1000000000000','Москва','40702810000000000001','Банк','044525000','30101810000000000002')
            conn.execute('INSERT OR REPLACE INTO portal_company_requisites(company_id,id,legal_name,inn,kpp,ogrn,legal_address,settlement_account,bank_name,bik,correspondent_account,updated_at) VALUES(1,1,?,?,?,?,?,?,?,?,?,?)',requisites+('2026-09-29',))
            conn.execute('INSERT OR REPLACE INTO portal_client_requisites(company_id,client_id,legal_name,inn,kpp,ogrn,legal_address,settlement_account,bank_name,bik,correspondent_account,updated_at) VALUES(1,1,?,?,?,?,?,?,?,?,?,?)',('ООО Клиент','7700000001','770001002','ОГРН 1000000000001','Москва','40702810000000000003','Банк клиента','044525001','30101810000000000004','2026-09-29'))
        work=self.work();invoice=self.post('invoices',dict(work_ids=[work['id']]))['data']
        generation=dict(document_type='invoice_xlsx',invoice_id=invoice['id'],request_id='invoice-xlsx-idempotent-'+uuid.uuid4().hex)
        doc=self.post('document-generate',generation)['data']
        repeated=self.post('document-generate',generation)['data']
        self.post('document-generate',dict(document_type='invoice_xlsx',invoice_id=invoice['id']),self.other_admin,status=400)
        self.assertEqual(doc['id'],repeated['id']);self.assertEqual(doc['category'],'invoice')
        self.assertEqual(doc['revision'],1);self.assertEqual(doc['mime_type'],'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        data=base64.b64decode(self.get('document-file?id='+doc['id'])['data']['file_b64'])
        book=load_workbook(BytesIO(data),data_only=False);sheet=book.active
        self.assertEqual(sheet['A1'].value,'Счёт на оплату')
        self.assertTrue(any(sheet.cell(row,1).value=='Операция' for row in range(1,sheet.max_row+1)))
        self.assertTrue(any([sheet.cell(row,col).value for col in range(1,5)]==['Операция','Количество','Цена','Сумма'] for row in range(1,sheet.max_row+1)))
        self.assertTrue(any(sheet.cell(row,1).value=='Исполнитель' for row in range(1,sheet.max_row+1)))
        self.assertTrue(any(sheet.cell(row,1).value=='Клиент' for row in range(1,sheet.max_row+1)))
        self.assertEqual(str(sheet.page_setup.paperSize),sheet.PAPERSIZE_A4)
        self.assertTrue(sheet.sheet_properties.pageSetUpPr.fitToPage)
        self.assertTrue(any(cell.data_type=='f' for row in sheet.iter_rows() for cell in row))

    def test_payroll_slip_uses_closed_snapshot_permissions_scope_and_immutable_facts(self):
        work=self.work();today=datetime.utcnow().date()
        if today.day>15:start=today.replace(day=1).isoformat();end=today.replace(day=15).isoformat()
        else:
            previous=today.replace(day=1)-timedelta(days=1)
            start=previous.replace(day=16).isoformat();end=previous.isoformat()
        with portal.db() as conn:
            repo=Repository(conn,1)
            source=dict(work,id=str(uuid.uuid4()),completed_at=end+'T12:00:00.000000',created_at=end+'T12:00:00.000000')
            repo.insert('works',{k:v for k,v in source.items() if k not in {'id','company_id'}},source['id'])
            conn.commit()
        period=self.post('payroll-periods',dict(period_start=start,period_end=end))['data']
        employee=period['snapshot']['employees'][0]
        with portal.db() as conn:
            payroll_employee=Repository(conn,1).payroll_employee(employee['employee_id'],legacy=True)['employee_id']
            settlements_before=Repository(conn,1).payroll_settlements(period['id'],payroll_employee)
        renderer=SimpleNamespace(payroll_slip_pdf=lambda *args: PDF)
        with patch.dict(sys.modules,{'pdf_documents':renderer}):
            doc=self.post('document-generate',dict(document_type='payroll_slip_pdf',payroll_period_id=period['id'],employee_id=payroll_employee))['data']
            self.post('document-generate',dict(document_type='payroll_slip_pdf',payroll_period_id=period['id'],employee_id=payroll_employee),self.worker,status=403)
            self.post('document-generate',dict(document_type='payroll_slip_pdf',payroll_period_id=period['id'],employee_id=payroll_employee),self.other_admin,status=400)
        xlsx_generation=dict(document_type='payroll_slip_xlsx',payroll_period_id=period['id'],employee_id=payroll_employee,
                             request_id='payroll-slip-xlsx-idempotent')
        xlsx=self.post('document-generate',xlsx_generation)['data']
        xlsx_repeat=self.post('document-generate',xlsx_generation)['data']
        self.post('document-generate',dict(document_type='payroll_slip_xlsx',payroll_period_id=period['id'],employee_id=payroll_employee),self.worker,status=403)
        self.post('document-generate',dict(document_type='payroll_slip_xlsx',payroll_period_id=period['id'],employee_id=payroll_employee),self.other_admin,status=400)
        from io import BytesIO
        from openpyxl import load_workbook
        slip=load_workbook(BytesIO(base64.b64decode(self.get('document-file?id='+xlsx['id'])['data']['file_b64'])),data_only=False).active
        self.assertEqual(xlsx['category'],'payroll');self.assertEqual(xlsx['employee_id'],payroll_employee)
        self.assertEqual(xlsx['id'],xlsx_repeat['id']);self.assertEqual(xlsx['payroll_period_id'],period['id'])
        self.assertEqual(xlsx['revision'],1);self.assertEqual(xlsx['mime_type'],'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        labels=[slip.cell(row,1).value for row in range(1,slip.max_row+1)]
        signatures=['Управляющий компанией','Управляющий подразделением','Сотрудник']
        self.assertEqual([x for x in labels if x in signatures],signatures)
        self.assertFalse(any(v in ('Клиент','Операция') for row in slip.iter_rows() for v in (cell.value for cell in row)))
        self.assertEqual(slip.cell(1,1).value,'Расчётный лист');self.assertEqual(str(slip.page_setup.paperSize),slip.PAPERSIZE_A4)
        self.assertEqual(doc['document_type'],'payroll_slip_pdf')
        self.assertEqual(doc['employee_id'],payroll_employee)
        self.assertEqual(doc['payroll_period_id'],period['id'])
        self.assertEqual(base64.b64decode(self.get('document-file?id='+doc['id'])['data']['file_b64']),PDF)
        saved=next(p for p in self.get('payroll-periods')['data'] if p['id']==period['id'])
        self.assertEqual(saved['snapshot'],period['snapshot'])
        with portal.db() as conn:
            settlements_after=Repository(conn,1).payroll_settlements(period['id'],payroll_employee)
        self.assertEqual(settlements_after,settlements_before)

class BlobStorageTest(unittest.TestCase):
    def test_content_addressed_company_keys_atomic_retry_and_traversal(self):
        with tempfile.TemporaryDirectory() as root:
            store=LocalFileStorage(root);a=store.put(PDF,1);b=store.put(PDF,2)
            self.assertNotEqual(a,b);self.assertEqual(store.put(PDF,1),a)
            self.assertEqual(store.get(a,1),PDF)
            for key,cid in (('../x',1),(a,2),('1/'+'x'*64,1)):
                with self.assertRaises(PermissionError):store.get(key,cid)
            self.assertFalse(list(Path(root).rglob('.tmp-*')))
