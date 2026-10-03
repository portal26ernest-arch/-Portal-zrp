"""Synthetic workbook regression; no database, documents or production data."""
from io import BytesIO
import unittest
from openpyxl import load_workbook
from financial_xlsx import invoice_xlsx, payroll_slip_xlsx

class FinancialTextTest(unittest.TestCase):
    def profile(self, text):
        return {key:text for key in ('legal_name','inn','legal_address','settlement_account','bank_name','bik','email')}

    def test_invoice_untrusted_strings_remain_text_and_total_is_formula(self):
        for hostile in ('=1+1', '=HYPERLINK("https://example.invalid/","label")', '#REF!', '+1+1', '@SUM(1,2)', 'Обычная работа'):
            with self.subTest(input=hostile):
                invoice={'id':'fixture','created_at':'2026-10-03','amount':200,
                         'lines':[{'operation_name':hostile,'quantity':2,'client_rate':100,'amount':200}]}
                sheet=load_workbook(BytesIO(invoice_xlsx(self.profile(hostile),self.profile(hostile),invoice)),data_only=False).active
                text_cells=[c for row in sheet for c in row if c.value==hostile]
                self.assertGreaterEqual(len(text_cells),3)
                self.assertTrue(all(c.data_type=='s' for c in text_cells))
                formulas=[c.value for row in sheet for c in row if c.data_type=='f']
                self.assertEqual(len(formulas),1)
                self.assertTrue(formulas[0].startswith('=SUM(D'))
                self.assertTrue(any(c.value==2 and c.data_type=='n' for row in sheet for c in row))

    def test_payroll_text_is_literal_and_numbers_preserved(self):
        for hostile in ('=1+1','#REF!','Обычный сотрудник'):
            sheet=load_workbook(BytesIO(payroll_slip_xlsx({'name':hostile},
                {'status':'closed','snapshot':{'fixture':True},'period_start':'2026-09-01','period_end':'2026-09-15'},
                {'display_name':hostile},{'balance':12345},'2026-10-03')),data_only=False).active
            self.assertEqual((sheet['B2'].value,sheet['B2'].data_type),(hostile,'s'))
            self.assertEqual((sheet['B4'].value,sheet['B4'].data_type),(hostile,'s'))
            self.assertEqual((sheet['B5'].value,sheet['B5'].data_type),(123.45,'n'))
            self.assertFalse(any(c.data_type=='f' for row in sheet for c in row))

if __name__=='__main__':unittest.main()
