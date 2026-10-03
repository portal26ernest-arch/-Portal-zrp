"""Canonical workbook contract, standard-reader round trip and tenant prefill."""
import base64
import hashlib
import io
import unittest
from zipfile import ZipFile
from xml.etree import ElementTree as ET
from excel_template import workbook,SHEETS,TEMPLATE_VERSION
from portal_excel_workbook import parse_template
import test_documents_api as fixtures

class TemplateTest(unittest.TestCase):
    def test_exact_sheets_headers_marker_machine_mapping_and_determinism(self):
        payload=workbook();parsed=parse_template(payload)
        self.assertEqual(parsed['sheets'],list(SHEETS))
        self.assertEqual(parsed['template_version'],TEMPLATE_VERSION)
        self.assertEqual(payload,workbook())
        self.assertEqual(parsed['rows'],{name:[] for name in SHEETS})
        with ZipFile(io.BytesIO(payload)) as book:
            for name in book.namelist():
                if name.endswith(('.xml','.rels')):ET.fromstring(book.read(name))
        self.assertEqual(parsed['checksum_sha256'],hashlib.sha256(payload).hexdigest())

    def test_money_dates_text_and_formula_like_names_are_not_executed(self):
        values={'Клиенты':[dict(client_ref='new',name='=HYPERLINK("bad")',inn='0012345678',active=1)],
                'Операции_Тарифы':[dict(client_ref='new',name='Упаковка',employee_rate='0.10',client_rate='12.34',active=1,effective_from='2027-01-01T00:00:00.000000')]}
        result=parse_template(workbook(values))['rows']
        self.assertEqual(result['Клиенты'][0]['inn'],'0012345678')
        self.assertEqual(result['Операции_Тарифы'][0]['employee_rate'],'0.10')
        self.assertEqual(result['Операции_Тарифы'][0]['effective_from'],'2027-01-01T00:00:00.000000')

    def test_standard_reader_opens_and_resaves_template(self):
        import openpyxl
        payload=workbook({'Клиенты':[dict(client_ref='new',name='Тест',active=1)]})
        book=openpyxl.load_workbook(io.BytesIO(payload))
        self.assertEqual(book.sheetnames,list(SHEETS))
        keys=[cell.value for cell in book['Клиенты'][2]]
        self.assertEqual(book['Клиенты'].cell(row=4,column=keys.index('name')+1).value,'Тест')
        self.assertEqual(book['Клиенты'].freeze_panes,'A4')
        self.assertTrue(book['Клиенты'].row_dimensions[1].hidden)
        self.assertTrue(book['Клиенты'].row_dimensions[2].hidden)
        from openpyxl.utils import get_column_letter
        for sheet_name in SHEETS:
            sheet=book[sheet_name];keys=[cell.value for cell in sheet[2]]
            for index,key in enumerate(keys,1):
                if key in {'company_id','employee_ref','employee_id','user_id','client_ref','client_id','operation_id','material_ref','material_id','norm_id','effective_to'}:
                    self.assertTrue(sheet.column_dimensions[get_column_letter(index)].hidden,(sheet_name,key))
        visible_headers=[cell.value for cell in book['Выработка'][3] if not book['Выработка'].column_dimensions[get_column_letter(cell.column)].hidden]
        self.assertEqual(visible_headers,['Сотрудник','Клиент','Операция','Количество','Товар / комментарий'])
        out=io.BytesIO();book.save(out)
        self.assertEqual(parse_template(out.getvalue())['rows']['Клиенты'][0]['name'],'Тест')

    def test_user_can_fill_and_resave_every_template_tab(self):
        import openpyxl
        book=openpyxl.load_workbook(io.BytesIO(workbook()))
        samples={
            'Компания':{'company_id':'1','name':'PORTAL','director':'Директор'},
            'Сотрудники':{'employee_ref':'employee:new','full_name':'Новый сотрудник','profile_username':'new.user','role':'packer','active':'1'},
            'Клиенты':{'client_ref':'client:new','name':'Новый клиент','active':'1','inn':'0012345678'},
            'Операции_Тарифы':{'client_ref':'client:new','name':'Упаковка','employee_rate':'5.00','client_rate':'8.00','active':'1','effective_from':'2027-01-01T00:00:00Z'},
            'Материалы':{'name':'Коробка','unit':'шт','unit_cost':'2.50','min_stock':'10','active':'1'},
            'Приход_материалов':{'material_name':'Коробка','quantity':'100','unit_cost':'2.50','note':'Стартовый приход'},
            'Нормы_материалов':{'client_name':'Новый клиент','operation_name':'Упаковка','material_name':'Коробка','qty_per_unit':'0.5','active':'1'},
            'Выработка':{'employee_name':'Новый сотрудник','client_name':'Новый клиент','operation_name':'Упаковка','quantity':'2','product':'Тест'},
        }
        for sheet_name,values in samples.items():
            sheet=book[sheet_name]
            keys=[cell.value for cell in sheet[2]]
            for key,value in values.items():
                sheet.cell(row=4,column=keys.index(key)+1,value=value)
        out=io.BytesIO();book.save(out)
        parsed=parse_template(out.getvalue())['rows']
        self.assertEqual(parsed['Компания'][0]['director'],'Директор')
        self.assertEqual(parsed['Сотрудники'][0]['full_name'],'Новый сотрудник')
        self.assertEqual(parsed['Клиенты'][0]['inn'],'0012345678')
        self.assertEqual(parsed['Операции_Тарифы'][0]['employee_rate'],'5.00')
        self.assertEqual(parsed['Операции_Тарифы'][0]['client_rate'],'8.00')
        self.assertEqual(parsed['Материалы'][0]['name'],'Коробка')
        self.assertEqual(parsed['Приход_материалов'][0]['quantity'],'100')
        self.assertEqual(parsed['Нормы_материалов'][0]['qty_per_unit'],'0.5')
        self.assertEqual(parsed['Выработка'][0]['quantity'],'2')

class TemplateAPITest(unittest.TestCase):
    request=fixtures.DocumentAPITest.request
    tearDown=fixtures.DocumentAPITest.tearDown
    setUp=fixtures.DocumentAPITest.setUp
    role_token=fixtures.DocumentAPITest.role_token
    get=fixtures.DocumentAPITest.get
    post=fixtures.DocumentAPITest.post

    def test_prefill_identity_rates_requisites_company_isolation_no_credentials(self):
        a=parse_template(base64.b64decode(self.get('document-template')['data']['file_b64']))
        b=parse_template(base64.b64decode(self.get('document-template',self.other_admin)['data']['file_b64']))
        self.assertEqual(a['rows']['Компания'][0]['company_id'],'1')
        self.assertEqual(b['rows']['Компания'][0]['company_id'],str(self.other))
        self.assertEqual(a['rows']['Операции_Тарифы'][0]['employee_rate'],'2.00')
        self.assertEqual(b['rows']['Операции_Тарифы'][0]['employee_rate'],'9.00')
        self.assertNotIn('Foreign only',str(a))
        for forbidden in ('pin_hash','pin_salt','token','telegram_id','Owner-secret-canary-123'):
            self.assertNotIn(forbidden,str(a))
        manager=self.role_token('manager')
        self.get('document-template',manager,status=403)
        self.assertEqual(parse_template(base64.b64decode(self.get('document-template-blank',manager)['data']['file_b64']))['rows']['Клиенты'],[])

    def test_generated_template_can_be_registered_in_documents(self):
        result=self.post('document-template-blank',{})['data']
        self.assertEqual(result['document_type'],'import_template_xlsx')
        self.assertEqual(self.get('documents')['data'][0]['id'],result['id'])
