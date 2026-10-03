"""Printable deterministic workbooks for immutable invoice/payroll snapshots."""
from datetime import datetime
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.page import PageMargins

def _data_cell(sheet,row,col,value):
    """Untrusted text must stay text; application formula cells are explicit."""
    cell=sheet.cell(row,col,value)
    if isinstance(value,str):cell.data_type='s'
    return cell

def _base(title):
    book=Workbook();sheet=book.active;sheet.title=title
    book.properties.creator='PORTAL';book.properties.created=datetime(2000,1,1);book.properties.modified=datetime(2000,1,1)
    sheet.page_setup.orientation='portrait';sheet.page_setup.paperSize=sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth=1;sheet.page_setup.fitToHeight=0;sheet.sheet_properties.pageSetUpPr.fitToPage=True
    sheet.page_margins=PageMargins(left=.25,right=.25,top=.4,bottom=.4,header=.2,footer=.2)
    sheet.sheet_view.showGridLines=False
    return book,sheet

def _header(sheet,row,values):
    for col,value in enumerate(values,1):
        cell=sheet.cell(row,col,value);cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='24476B');cell.alignment=Alignment(wrap_text=True,vertical='center')

def _requisites(sheet,start,label,data):
    sheet.cell(start,1,label).font=Font(bold=True,size=12)
    fields=(('legal_name','Наименование'),('inn','ИНН'),('kpp','КПП'),('ogrn','ОГРН'),('legal_address','Адрес'),('settlement_account','Расчётный счёт'),('bank_name','Банк'),('bik','БИК'),('correspondent_account','Корреспондентский счёт'),('phone','Телефон'),('email','Email'),('tax_info','Налогообложение'))
    row=start+1
    for key,name in fields:
        value=data.get(key)
        if value not in (None,''):
            sheet.cell(row,1,name);_data_cell(sheet,row,2,str(value));row+=1
    return row

def _finish(book,sheet,widths):
    for col,width in enumerate(widths,1):sheet.column_dimensions[chr(64+col)].width=width
    sheet.freeze_panes='A1';sheet.print_options.horizontalCentered=True
    sheet.print_title_rows='1:1';sheet.sheet_properties.pageSetUpPr.fitToPage=True
    out=BytesIO();book.save(out);return out.getvalue()

def invoice_xlsx(company,client,invoice):
    for label,data in (('исполнителя',company),('клиента',client)):
        missing=[field for field in ('legal_name','inn','legal_address','settlement_account','bank_name','bik') if data.get(field) in (None,'')]
        if missing:raise ValueError(f'Не заполнены реквизиты {label}: {", ".join(missing)}')
    if not invoice.get('created_at') or not invoice.get('lines'):
        raise ValueError('В снимке счёта отсутствуют дата или строки')
    for line in invoice['lines']:
        if line.get('operation_name') in (None,'') or any(line.get(k) is None for k in ('quantity','client_rate','amount')):
            raise ValueError('В снимке счёта отсутствуют подтверждённые строки операции')
        if type(line['quantity']) is not int or line['quantity']<1 or type(line['client_rate']) is not int or line['client_rate']<0 or type(line['amount']) is not int or line['amount']<0:
            raise ValueError('В снимке счёта указаны некорректные количество или сумма')
        if line['quantity']*line['client_rate']!=line['amount']:
            raise ValueError('Сумма строки счёта не совпадает с количеством и ценой')
    if type(invoice.get('amount')) is not int or sum(line['amount'] for line in invoice['lines'])!=invoice['amount']:
        raise ValueError('Итог счёта не совпадает с суммой строк')
    book,sheet=_base('Счёт')
    number=str(invoice.get('number') or invoice.get('id') or '')
    issued=str(invoice.get('created_at',''))[:10]
    title='Счёт на оплату'+((' № '+number) if number else '')+((' от '+issued) if issued else '')
    sheet.append([title])
    sheet.merge_cells('A1:D1')
    sheet['A1'].font=Font(bold=True,size=18,color='24476B')
    row=_requisites(sheet,3,'Исполнитель',company);row=_requisites(sheet,row+1,'Клиент',client)
    sheet.cell(row,1,'Дата');_data_cell(sheet,row,2,str(invoice.get('created_at',''))[:10]);row+=1
    if invoice.get('due_at'):sheet.cell(row,1,'Оплатить до');_data_cell(sheet,row,2,str(invoice['due_at'])[:10]);row+=1
    row+=1;_header(sheet,row,['Операция','Количество','Цена','Сумма']);first=row+1
    for line in invoice['lines']:
        row+=1;_data_cell(sheet,row,1,line['operation_name']);sheet.cell(row,2,line['quantity']);sheet.cell(row,3,line['client_rate']/100);sheet.cell(row,4,line['amount']/100)
        sheet.cell(row,3).number_format=sheet.cell(row,4).number_format='#,##0.00'
    row+=1;sheet.cell(row,3,'Итого, RUB').font=Font(bold=True);sheet.cell(row,4,f'=SUM(D{first}:D{row-1})').font=Font(bold=True);sheet.cell(row,4).number_format='#,##0.00'
    sheet.print_area=f'A1:D{row}';sheet.auto_filter.ref=f'A{row-len(invoice["lines"])}:D{row-1}'
    return _finish(book,sheet,[42,16,18,20])

def payroll_slip_xlsx(company,period,employee,summary,issue_date):
    if period.get('status')!='closed' or not period.get('snapshot'):
        raise ValueError('Расчётный лист доступен только по закрытому снимку периода')
    if not (company.get('name') or company.get('legal_name')) or not employee.get('display_name') or not issue_date:
        raise ValueError('Для расчётного листа отсутствуют обязательные данные')
    book,sheet=_base('Расчётный лист');sheet.append(['Расчётный лист']);sheet['A1'].font=Font(bold=True,size=18,color='24476B')
    rows=[('Компания',company.get('name') or company.get('legal_name')),('Период / год',f'{period["period_start"]} — {period["period_end"]} / {period["period_start"][:4]}'),('ФИО',employee.get('display_name')),('Итоговая сумма, RUB',summary['balance']/100),('Дата',issue_date),('Управляющий компанией','________________________'),('Управляющий подразделением','________________________'),('Сотрудник','________________________')]
    for row,values in enumerate(rows,2):
        for col,value in enumerate(values,1):_data_cell(sheet,row,col,value)
    sheet['B5'].number_format='#,##0.00';sheet.print_area='A1:B9'
    return _finish(book,sheet,[38,44])
