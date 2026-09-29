"""Contract tests for payroll workbook payment-sheet export."""
from io import BytesIO
import unittest

from openpyxl import load_workbook

from report_xlsx import payroll_xlsx


class PayrollPaymentSheetTest(unittest.TestCase):
    def test_payments_sheet_reports_ledger_totals_dates_without_changing_other_sheets(self):
        snapshot={
            'period_start':'2026-09-01','period_end':'2026-09-15','total_quantity':3,
            'total_salary':7500,'work_rows':2,
            'employees':[{'display_name':'Synthetic Employee','employee_id':7,'user_id':9,
                          'quantity':3,'salary':7500,'work_rows':2}],
            'details':[],
        }
        settlements={
            'employees':[{'display_name':'Synthetic Employee','employee_id':7,
                          'accrued':7500,'paid':2500,'balance':5000}],
            'totals':{'accrued':7500,'paid':2500,'balance':5000},
            'entries':[
                {'employee_id':7,'effect':'payment','amount':1500,'occurred_at':'2026-09-04T10:30:00Z'},
                {'employee_id':7,'effect':'payment','amount':1000,'occurred_at':'2026-09-10T12:00:00Z'},
                {'employee_id':7,'effect':'entitlement','amount':500,'occurred_at':'2026-09-12T12:00:00Z'},
            ],
        }
        book=load_workbook(BytesIO(payroll_xlsx(snapshot,settlements)),data_only=False)
        self.assertEqual(book.sheetnames[:4],['Сводка','Выплаты','Сотрудники','Детализация'])
        sheet=book['Выплаты']
        self.assertEqual([sheet.cell(1,col).value for col in range(1,7)],
                         ['Сотрудник','Период','Начислено, ₽','Выплаты, ₽','Остаток, ₽','Даты выплат'])
        self.assertEqual([sheet.cell(2,col).value for col in range(1,7)],
                         ['Synthetic Employee','2026-09-01 — 2026-09-15',75.0,25.0,50.0,
                          '2026-09-04, 2026-09-10'])
        self.assertEqual([sheet.cell(3,col).value for col in range(1,6)],
                         ['Итого','2026-09-01 — 2026-09-15',75.0,25.0,50.0])
        self.assertEqual(book['Сводка']['A4'].value,'Начислено, ₽')
        self.assertEqual(book['Сотрудники']['A2'].value,'Synthetic Employee')
        self.assertEqual(book['Детализация'].max_row,1)

    def test_no_settlement_data_exports_honest_empty_payment_rows(self):
        snapshot={'period_start':'2026-09-01','period_end':'2026-09-15','total_quantity':0,
                  'total_salary':0,'work_rows':0,'employees':[],'details':[]}
        book=load_workbook(BytesIO(payroll_xlsx(snapshot)),data_only=False)
        self.assertIn('Выплаты',book.sheetnames)
        sheet=book['Выплаты']
        self.assertEqual(sheet.max_row,2)
        self.assertEqual([sheet.cell(2,col).value for col in range(1,6)],
                         ['Итого','2026-09-01 — 2026-09-15',0.0,0.0,0.0])


if __name__=='__main__':
    unittest.main()
